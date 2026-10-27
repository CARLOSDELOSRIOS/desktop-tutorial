"""Renderiza fuentes editables a DOCX/PDF y PPTX/PDF; no crea resultados SQL."""
from pathlib import Path
import json,subprocess,shutil,io,re,html,textwrap
from docx import Document
from docx.shared import Inches as DInches,Pt as DPt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table as DTable
from docx.text.paragraph import Paragraph as DParagraph
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,TableStyle,Image as RImage,PageBreak,Preformatted,KeepTogether
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
NAVY='#123d51';TEAL='#147d92';INK='#233a47';GOLD='#e5b95c'

def fonts():
    paths=[Path('/usr/share/fonts/truetype/dejavu'),Path('C:/Windows/Fonts')]
    for folder in paths:
        normal=folder/('DejaVuSans.ttf' if 'dejavu' in str(folder) else 'arial.ttf')
        bold=folder/('DejaVuSans-Bold.ttf' if 'dejavu' in str(folder) else 'arialbd.ttf')
        mono=folder/('DejaVuSansMono.ttf' if 'dejavu' in str(folder) else 'consola.ttf')
        if normal.exists() and bold.exists() and mono.exists():
            pdfmetrics.registerFont(TTFont('Body',str(normal)));pdfmetrics.registerFont(TTFont('BodyBold',str(bold)));pdfmetrics.registerFont(TTFont('Mono',str(mono)));return
    raise RuntimeError('No se encontraron fuentes Arial/Consolas o DejaVu Sans para PDF')

def make_docx():
    if not shutil.which('pandoc'):raise RuntimeError('Instalar Pandoc para convertir el informe Markdown a DOCX')
    subprocess.run(['pandoc','informe.md','-f','markdown','-t','docx','--standalone','-o','informe.docx'],cwd=ROOT/'reports',check=True)
    path=ROOT/'reports/informe.docx';doc=Document(path)
    for sec in doc.sections:
        sec.page_width=DInches(8.27);sec.page_height=DInches(11.69);sec.left_margin=DInches(.7);sec.right_margin=DInches(.7);sec.top_margin=DInches(.65);sec.bottom_margin=DInches(.65)
    style_by_name={s.name:s for s in doc.styles}
    for name in ('Normal','Body Text','First Paragraph'):
        style=style_by_name[name];style.font.name='Arial';style.font.size=DPt(10);style.paragraph_format.space_after=DPt(6)
    for name in ('Heading 1','Heading 2','Heading 3'):
        style_by_name[name].font.name='Arial';style_by_name[name].font.color.rgb=__import__('docx').shared.RGBColor.from_string('123D51');style_by_name[name].paragraph_format.keep_with_next=True
    if 'Source Code' in style_by_name:
        style_by_name['Source Code'].font.name='Consolas';style_by_name['Source Code'].font.size=DPt(7)
    for t in doc.tables:
        t.autofit=True
        pr=t.rows[0]._tr.get_or_add_trPr();flag=OxmlElement('w:tblHeader');pr.append(flag)
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=DPt(2)
                    for r in p.runs:r.font.name='Arial';r.font.size=DPt(7.5)
    for shape in doc.inline_shapes:
        if shape.width>DInches(6.8):
            ratio=DInches(6.8)/shape.width;shape.width=DInches(6.8);shape.height=int(shape.height*ratio)
    for p in doc.paragraphs:
        if p.style.name=='Heading 1' and (p.text.startswith('1. ') or p.text.startswith('Anexo ')):p.paragraph_format.page_break_before=True
    footer=doc.sections[0].footer.paragraphs[0];footer.text='PAC PSMA00611 · Datos sintéticos · Página '
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    doc.save(path);return path

def report_pdf(docx_path):
    doc=Document(docx_path);styles=getSampleStyleSheet();width=A4[0]-88
    body=ParagraphStyle('BodyPDF',fontName='Body',fontSize=9.6,leading=14,spaceAfter=7,textColor=colors.HexColor(INK))
    heading=ParagraphStyle('H1PDF',parent=body,fontName='BodyBold',fontSize=17,leading=21,spaceBefore=14,spaceAfter=10,keepWithNext=True,textColor=colors.HexColor(NAVY))
    h2=ParagraphStyle('H2PDF',parent=heading,fontSize=12,leading=16,spaceBefore=9)
    small=ParagraphStyle('SmallPDF',parent=body,fontSize=7.1,leading=9.4,spaceAfter=2)
    code=ParagraphStyle('CodePDF',fontName='Mono',fontSize=6.2,leading=8.4,spaceAfter=9)
    story=[]
    def esc(s):return html.escape(s).replace('\n','<br/>')
    for elem in doc.element.body.iterchildren():
        if elem.tag==qn('w:p'):
            p=DParagraph(elem,doc);name=p.style.name;txt=p.text
            for blip in p._p.xpath('.//a:blip'):
                blob=doc.part.related_parts[blip.get(qn('r:embed'))].blob
                im=Image.open(io.BytesIO(blob));iw,ih=im.size;scale=min(width/iw,360/ih)
                story += [Spacer(1,7),RImage(io.BytesIO(blob),width=iw*scale,height=ih*scale),Spacer(1,7)]
            if not txt.strip():continue
            if name.startswith('Heading'):
                if txt.startswith('1. ') or txt.startswith('Anexo '):story.append(PageBreak())
                story.append(Paragraph(esc(txt),heading if name=='Heading 1' else h2))
            elif name=='Source Code':story.append(Preformatted(txt,code,maxLineLength=120))
            else:story.append(Paragraph(esc(txt),body))
        elif elem.tag==qn('w:tbl'):
            table=DTable(elem,doc);data=[[Paragraph(esc(c.text),small) for c in row.cells] for row in table.rows]
            n=len(table.columns)
            if n==5:widths=[width*x for x in [.17,.23,.06,.14,.40]]
            elif n==3:widths=[width*.10,width*.56,width*.34]
            elif n==4:widths=[width*.12,width*.28,width*.30,width*.30]
            else:widths=[width/n]*n
            t=LongTable(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#deedf2')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#adc5ce')),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
            story += [t,Spacer(1,10)]
    def page(c,doc):
        c.setStrokeColor(colors.HexColor(TEAL));c.line(44,A4[1]-26,A4[0]-44,A4[1]-26)
        c.setFont('Body',7);c.setFillColor(colors.HexColor(INK));c.drawString(44,24,'PAC PSMA00611 · Datos sintéticos · SQL/Azure pendientes');c.drawRightString(A4[0]-44,24,str(doc.page))
    SimpleDocTemplate(str(ROOT/'reports/informe.pdf'),pagesize=A4,rightMargin=44,leftMargin=44,topMargin=40,bottomMargin=42,title='Base de datos de aerolínea — PAC PSMA00611',author='[POR COMPLETAR]').build(story,onFirstPage=page,onLaterPages=page)

def textbox(slide,x,y,w,h,text,size,color,bold=False):
    tf=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
    p=tf.paragraphs[0];p.text=text;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color.lstrip('#'));return tf

def make_slides():
    data=json.loads((ROOT/'slides/content.json').read_text(encoding='utf8'));assert 10<=len(data)<=15 and sum(x['seconds'] for x in data)==1200
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    pdf=canvas.Canvas(str(ROOT/'slides/presentacion.pdf'),pagesize=(960,540));pdf.setTitle('PAC PSMA00611 — Presentación de 20 minutos')
    for i,s in enumerate(data,1):
        sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string('123D51')
        textbox(sl,.6,.35,12.1,.38,'PSMA00611  /  BASES DE DATOS',12,'#9ccbd7')
        textbox(sl,.6,.95,12.1,.85,s['title'],32,'#ffffff',True)
        if i==1:textbox(sl,.65,1.85,11.9,.6,s['subtitle'],22,'#e5b95c')
        pdf.setFillColor(colors.HexColor(NAVY));pdf.rect(0,0,960,540,fill=1,stroke=0)
        pdf.setFillColor(colors.HexColor('#9ccbd7'));pdf.setFont('Body',10);pdf.drawString(43,504,'PSMA00611  /  BASES DE DATOS')
        pdf.setFillColor(colors.white);pdf.setFont('BodyBold',27);pdf.drawString(43,452,s['title'])
        if i==1:pdf.setFont('Body',18);pdf.setFillColor(colors.HexColor(GOLD));pdf.drawString(47,395,s['subtitle'])
        if s.get('image'):
            path=ROOT/s['image'];im=Image.open(path);iw,ih=im.size
            # Una placa blanca conserva el contraste de diagramas vectoriales exportados.
            from pptx.enum.shapes import MSO_SHAPE
            rect=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(.65),Inches(1.95),Inches(12),Inches(3.75));rect.fill.solid();rect.fill.fore_color.rgb=RGBColor(255,255,255);rect.line.fill.background()
            scale=min(11.5/iw,3.4/ih);w,h=iw*scale,ih*scale;sl.shapes.add_picture(str(path),Inches((13.333-w)/2),Inches(2.12+(3.4-h)/2),width=Inches(w),height=Inches(h))
            pdf.setFillColor(colors.white);pdf.roundRect(47,129,864,270,8,fill=1,stroke=0)
            scale_pdf=min(828/iw,245/ih);w2,h2=iw*scale_pdf,ih*scale_pdf;pdf.drawImage(str(path),(960-w2)/2,141+(245-h2)/2,width=w2,height=h2)
            y=5.95;pdf_y=102;size=18
        else:y=2.7 if i==1 else 2.15;pdf_y=332 if i==1 else 365;size=25
        for bullet in s['bullets']:
            tf=textbox(sl,.75,y,11.9,.94,'•  '+bullet,size,'#ffffff');tf.paragraphs[0].space_after=Pt(9)
            lines=textwrap.wrap('•  '+bullet,width=87 if size==18 else 69)
            pdf.setFillColor(colors.white);pdf.setFont('Body',13.5 if size==18 else 18)
            for line in lines:pdf.drawString(53,pdf_y,line);pdf_y-=20 if size==18 else 27
            y+=.50 if size==18 else 1.03;pdf_y-=9
        textbox(sl,.65,7.02,10,.25,f'Datos sintéticos · Tiempo de esta lámina: {s["seconds"]} s',10,'#9ccbd7');textbox(sl,12.0,7.0,.7,.3,f'{i:02}',12,'#e5b95c')
        sl.notes_slide.notes_text_frame.text=s['notes']+f'\nTiempo: {s["seconds"]} segundos. Total de exposición: 1200 segundos.'
        pdf.setFillColor(colors.HexColor('#9ccbd7'));pdf.setFont('Body',9);pdf.drawString(47,21,f'Datos sintéticos · Tiempo de esta lámina: {s["seconds"]} s');pdf.drawRightString(913,21,f'{i:02} / {len(data)}');pdf.showPage()
    prs.save(ROOT/'slides/presentacion.pptx');pdf.save()
    return len(data)

if __name__=='__main__':
    fonts();path=make_docx();report_pdf(path);n=make_slides();print(f'Informe DOCX/PDF y {n} diapositivas PPTX/PDF generados; resultados SQL no fabricados.')
